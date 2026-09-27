import json
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from PIL import Image

from skfs_ocr.config import ROOT, load_config
from skfs_ocr.dates import guess_month, parse_file_date, parse_page_date, resolve_date
from skfs_ocr.pipeline import cache_file, load_parties, run
from skfs_ocr.validate import build_day, check_nozzle, merge_extractions, month_checks
from skfs_ocr.workbook import day_cell_values

OCT = Path("FY2025-26") / "07 Oct-2025"
NOV = Path("FY2025-26") / "08 Nov-2025"

FIX = Path(__file__).parent / "fixtures"


def fixture(name):
    return json.loads((FIX / name).read_text(encoding="utf-8"))["extraction"]


@pytest.fixture
def cfg(tmp_path):
    c = load_config()
    for key in ("input", "raw", "output", "extracted"):
        c["paths"][key] = tmp_path / key
        c["paths"][key].mkdir()
    return c


def add_photo(cfg, name, extraction, color=(200, 200, 200)):
    """A dummy photo plus its already-read JSON (so no API call happens)."""
    from skfs_ocr.extract import save_record, sha256

    photo = cfg["paths"]["input"] / name
    photo.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (40, 30), color).save(photo)
    save_record(cache_file(cfg, photo), {"source": photo.name, "sha256": sha256(photo),
                                         "read_by": "test", "extraction": extraction})
    return photo


# ---- dates ----------------------------------------------------------------

@pytest.mark.parametrize("stem,expected", [
    ("03", date(2025, 10, 3)),
    ("12.10", date(2025, 10, 12)),
    ("30-10-2025", date(2025, 10, 30)),
    ("2025-10-07", date(2025, 10, 7)),
    ("07 (2)", date(2025, 10, 7)),
    ("7a", date(2025, 10, 7)),
    ("IMG_20251007", None),
])
def test_file_dates(stem, expected):
    assert parse_file_date(stem, 2025, 10) == expected


def test_page_dates():
    assert parse_page_date("3/10/25", None, None) == date(2025, 10, 3)
    assert parse_page_date("30|11|25", None, None) == date(2025, 11, 30)
    assert parse_page_date("31/11/25", None, None) is None   # not a real date


def test_month_guess_from_pages_and_names():
    items = [("01.10", {"date_text": "1/10/25"}), ("03", {"date_text": "3/10/25"}),
             ("04", {"date_text": None})]
    assert guess_month(items) == (2025, 10)


def test_name_and_page_disagree_warns():
    d, errors, warnings = resolve_date("05", {"date_text": "6/10/25"}, 2025, 10)
    assert d == date(2025, 10, 5) and not errors and "using file name" in warnings[0]


# ---- checks ----------------------------------------------------------------

def test_clean_day_has_no_errors(cfg):
    day = build_day(date(2025, 10, 3), ["03.jpeg"], fixture("day_clean.json"), cfg, load_parties(cfg))
    assert day.errors == []
    assert day.fuel_totals == {"petrol": 161380.0, "diesel": 181023.0, "power": 15833.0}
    assert sum(day.fuel_totals.values()) == 358236


def test_misread_digit_is_caught(cfg):
    ext = fixture("day_clean.json")
    ext["nozzles"][1]["final_amount"] = 75924     # 75024 misread as 75924
    day = build_day(date(2025, 10, 3), ["03.jpeg"], ext, cfg, load_parties(cfg))
    assert any("total sale written" in e for e in day.errors)


def test_payment_misread_is_caught(cfg):
    ext = fixture("day_clean.json")
    ext["payments"]["cash"] = 75777
    day = build_day(date(2025, 10, 3), ["03.jpeg"], ext, cfg, load_parties(cfg))
    assert any("payment lines add up" in e for e in day.errors)


def test_nozzle_arithmetic():
    rates = {"petrol": 104.71, "diesel": 90.21, "power": 112.80}
    tol = load_config()["tolerance"]
    nz = {"label": "MS A2", "opening": 1335159.50, "closing": 1335989.57, "difference": 830.07,
          "testing": 5, "net_litres": 825.07, "rate": 104.71, "amount_written": 86393,
          "adjustments": [-37], "final_amount": 86356}
    assert check_nozzle(nz, rates, tol) == ("petrol", 86356, [], [])
    nz["closing"] = 1335999.57
    assert "closing" in check_nozzle(nz, rates, tol)[3][0]


def test_wrong_rate_written_uses_amount(cfg):
    rates = {"petrol": 104.71, "diesel": 90.21, "power": 112.80}
    nz = {"label": "HSD A1", "opening": None, "closing": None, "difference": 581.03,
          "testing": 5, "net_litres": 576.03, "rate": 104.71, "amount_written": 51963,
          "adjustments": [], "final_amount": 51963}
    fuel, amount, errs, _ = check_nozzle(nz, rates, cfg["tolerance"])
    assert (fuel, amount, errs) == ("diesel", 51963, [])


def test_meter_continuity(cfg):
    book = load_parties(cfg)
    a = build_day(date(2025, 10, 3), [], fixture("day_clean.json"), cfg, book)
    nxt = fixture("day_mismatch.json")
    b = build_day(date(2025, 10, 4), [], nxt, cfg, book)
    month_checks([a, b])
    assert not any("previous day closed" in w for w in b.warnings)
    nxt["nozzles"][0]["opening"] = 105738.40
    b = build_day(date(2025, 10, 4), [], nxt, cfg, book)
    month_checks([a, b])
    assert any("previous day closed" in w for w in b.warnings)


def test_party_names_normalised(cfg):
    book = load_parties(cfg)
    assert book.match("Mool Chand")["name"] == "mool Chand yadav"
    assert book.match("Shyam Transport")["name"] == "Shyam Tran."
    assert book.match("piploda ")["name"] == "Piploda"
    assert book.match("Zzz Unknown") is None


def test_two_photos_same_day_merge():
    left = {"rates": {"petrol": 104.71}, "payments": {}, "nozzles": [{"label": "MS"}]}
    right = {"rates": {"petrol": 104.71}, "payments": {"cash": 5}, "credit_entries": [{"amount": 1}]}
    merged, errors = merge_extractions([("a", left), ("b", right)])
    assert not errors and merged["payments"]["cash"] == 5 and len(merged["nozzles"]) == 1
    right["rates"]["petrol"] = 105
    assert merge_extractions([("a", left), ("b", right)])[1]


# ---- workbook ----------------------------------------------------------------

def test_cells_match_sample_layout(cfg):
    day = build_day(date(2025, 10, 3), [], fixture("day_clean.json"), cfg, load_parties(cfg))
    v = day_cell_values(cfg, day)
    assert (v["C4"], v["C5"], v["C6"]) == (104.71, 90.21, 112.80)
    assert v["H13"] + v["H14"] == 161380 and v["H15"] is None
    assert v["I13"] + v["I14"] + v["I15"] == 181023
    assert v["J13"] == 15833
    assert v["T8"] == 90869 and v["R8"] == 75717 and v["L8"] == 3673.71
    assert v["N8"] == round(15815.65 + 95681.60, 2)
    assert v["A15"] == "mool Chand yadav" and v["F15"] == 22343 and v["E15"] == 20454.21
    assert v["B15"] == "diesal"
    assert "K17" not in v and "K15" not in v           # old manual cells no longer written


def test_full_run_writes_verified_workbook(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    add_photo(cfg, "04.jpeg", fixture("day_mismatch.json"), color=(10, 10, 10))
    code = run(cfg, use_api=False, log=lambda *a: None)
    assert code == 1                          # day 4 has a real ledger mismatch
    out = cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx"
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["3-10-2025", "4-10-2025"]
    ws = wb["3-10-2025"]
    tpl = openpyxl.load_workbook(cfg["paths"]["template"])["TEMPLATE"]
    # every formula of the template survives unchanged
    for row in tpl.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert ws[c.coordinate].value == c.value, c.coordinate
    assert ws["C4"].value == 104.71 and ws["R8"].value == 75717
    report = (cfg["paths"]["output"] / OCT / "sale Oct-2025 - check report.txt").read_text(encoding="utf-8")
    assert "03-10-2025  [OK]" in report and "04-10-2025  [ERROR]" in report
    assert "No photo for day(s): 1, 2, 5," in report


def test_rerun_is_identical_and_needs_no_api(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    run(cfg, use_api=False, log=lambda *a: None)
    out = cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx"
    first = {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in openpyxl.load_workbook(out)}
    assert run(cfg, use_api=False, log=lambda *a: None) == 0
    second = {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in openpyxl.load_workbook(out)}
    assert first == second


def test_new_photo_without_key_gets_blank_json_and_is_reported(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    new = cfg["paths"]["input"] / "05.jpeg"
    Image.new("RGB", (40, 30), (1, 2, 3)).save(new)
    logs = []
    assert run(cfg, use_api=False, log=logs.append) == 1
    assert cache_file(cfg, new).exists()
    assert any("05.jpeg: not read yet" in m for m in logs)


def test_broken_json_is_reported_not_crashing(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    bad = add_photo(cfg, "06.jpeg", fixture("day_clean.json"), color=(9, 9, 9))
    cache_file(cfg, bad).write_text("{ not json", encoding="utf-8")
    logs = []
    assert run(cfg, use_api=False, log=logs.append) == 1
    assert any("06.jpeg" in m and "not valid JSON" in m for m in logs)


def test_month_folders_make_separate_workbooks(cfg):
    add_photo(cfg, "2025-10/03.jpeg", fixture("day_clean.json"))
    nov = fixture("day_clean.json") | {"date_text": "3/11/25"}
    add_photo(cfg, "2025-11/03.jpeg", nov, color=(50, 60, 70))
    run(cfg, use_api=False, log=lambda *a: None)
    assert (cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx").exists()
    assert (cfg["paths"]["output"] / NOV / "sale Nov-2025.xlsx").exists()


def test_template_is_the_committed_one():
    assert (ROOT / "template" / "template.xlsx").exists()


def test_loose_photos_of_two_months_split(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    add_photo(cfg, "04.jpeg", fixture("day_mismatch.json"), color=(1, 1, 1))
    add_photo(cfg, "03.11.jpeg", fixture("day_clean.json") | {"date_text": "3/11/25"}, color=(5, 5, 5))
    run(cfg, use_api=False, log=lambda *a: None)
    oct_ = openpyxl.load_workbook(cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx")
    nov = openpyxl.load_workbook(cfg["paths"]["output"] / NOV / "sale Nov-2025.xlsx")
    assert oct_.sheetnames[:2] == ["3-10-2025", "4-10-2025"] and nov.sheetnames[0] == "3-11-2025"


# ---- filing photos by financial year ------------------------------------------

def test_fy_folders():
    from skfs_ocr.dates import parse_month_folder
    from skfs_ocr.workbook import month_dir

    c = load_config()
    assert month_dir(c, 2025, 4) == Path("FY2025-26/01 Apr-2025")
    assert month_dir(c, 2025, 10) == Path("FY2025-26/07 Oct-2025")
    assert month_dir(c, 2026, 3) == Path("FY2025-26/12 Mar-2026")
    assert month_dir(c, 2026, 4) == Path("FY2026-27/01 Apr-2026")
    assert parse_month_folder("12 Mar-2026") == (2026, 3)
    assert parse_month_folder("2025-11") == (2025, 11)


def test_photos_move_to_raw_data_and_input_is_cleared(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    add_photo(cfg, "sub/04.jpeg", fixture("day_mismatch.json"), color=(1, 1, 1))
    new = cfg["paths"]["input"] / "05.jpeg"                 # not read yet -> stays
    Image.new("RGB", (40, 30), (1, 2, 3)).save(new)
    run(cfg, use_api=False, log=lambda *a: None)
    raw = cfg["paths"]["raw"] / OCT
    assert sorted(p.name for p in raw.iterdir()) == ["03.jpeg", "04.jpeg"]
    assert sorted(p.name for p in cfg["paths"]["input"].rglob("*")) == ["05.jpeg"]
    assert (cfg["paths"]["extracted"] / OCT / "03.jpeg.json").exists()
    assert not (cfg["paths"]["extracted"] / "03.jpeg.json").exists()

    # next run: input has only the unread photo, month still has both days
    run(cfg, use_api=False, log=lambda *a: None)
    wb = openpyxl.load_workbook(cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx")
    assert wb.sheetnames[:2] == ["3-10-2025", "4-10-2025"]


def test_adding_a_day_later_keeps_earlier_days(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    run(cfg, use_api=False, log=lambda *a: None)
    assert not list(cfg["paths"]["input"].iterdir())
    later = fixture("day_mismatch.json")
    add_photo(cfg, "04.jpeg", later, color=(1, 1, 1))
    run(cfg, use_api=False, log=lambda *a: None)
    wb = openpyxl.load_workbook(cfg["paths"]["output"] / OCT / "sale Oct-2025.xlsx")
    assert wb.sheetnames[:2] == ["3-10-2025", "4-10-2025"]
    assert sorted(p.name for p in (cfg["paths"]["raw"] / OCT).iterdir()) == ["03.jpeg", "04.jpeg"]


def test_same_name_different_photo_is_kept_both(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    run(cfg, use_api=False, log=lambda *a: None)
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"), color=(7, 7, 7))   # second page of day 3
    run(cfg, use_api=False, log=lambda *a: None)
    names = sorted(p.name for p in (cfg["paths"]["raw"] / OCT).iterdir())
    assert names == ["03 (2).jpeg", "03.jpeg"]


def test_same_photo_put_in_twice_is_not_duplicated(cfg):
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))
    run(cfg, use_api=False, log=lambda *a: None)
    add_photo(cfg, "03.jpeg", fixture("day_clean.json"))                    # identical file
    run(cfg, use_api=False, log=lambda *a: None)
    assert [p.name for p in (cfg["paths"]["raw"] / OCT).iterdir()] == ["03.jpeg"]
    assert not list(cfg["paths"]["input"].iterdir())


def test_cash_qty_matches_owner_workbook(cfg):
    """Q4:Q6 = sale qty - Paytm/ICICI/PhonePe/HP/udhar qty, as in the owner's final October file."""
    from skfs_ocr.config import ROOT

    ref = openpyxl.load_workbook(ROOT / "template" / "sample sale Oct-2025.xlsx")["3-10-2025"]
    day = build_day(date(2025, 10, 3), [], fixture("day_clean.json"), cfg, load_parties(cfg))
    v = day_cell_values(cfg, day)
    for cell in ("Q4", "Q5", "Q6"):
        assert v[cell] == pytest.approx(ref[cell].value, abs=1e-9)


def test_template_keeps_owner_formulas():
    from skfs_ocr.config import ROOT

    ref = openpyxl.load_workbook(ROOT / "template" / "sample sale Oct-2025.xlsx")["1-10-2025"]
    tpl = openpyxl.load_workbook(ROOT / "template" / "template.xlsx")["TEMPLATE"]
    for row in ref.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert tpl[c.coordinate].value == c.value, c.coordinate
