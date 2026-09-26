"""Write a month of Day objects into a copy of the template workbook.

Every day sheet is a copy of the TEMPLATE sheet, so all formulas, labels,
widths and formats stay exactly as in your own workbook; only the input
cells listed in config.yaml are written. After saving, the file is opened
again and every written cell is compared with what should be there.
"""
import os
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl

from .template_tool import day_input_cells
from .validate import Day

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class WorkbookError(Exception):
    pass


def month_dir(cfg: dict, year: int, month: int) -> Path:
    """FY2025-26/07 Oct-2025 (financial year April-March)."""
    from .dates import fy_of

    wb, fy = cfg["workbook"], fy_of(year, month)
    fields = dict(year=year, month=month, mon=MONTHS[month - 1], fy_start=fy,
                  fy_end=(fy + 1) % 100, fy_month=(month - 4) % 12 + 1)
    return Path(wb["fy_folder"].format(**fields)) / wb["month_folder"].format(**fields)


def output_path(cfg: dict, year: int, month: int) -> Path:
    name = cfg["workbook"]["file_name"].format(year=year, month=month, mon=MONTHS[month - 1])
    return cfg["paths"]["output"] / month_dir(cfg, year, month) / name


def sheet_name(cfg: dict, d) -> str:
    return cfg["workbook"]["day_sheet_name"].format(
        d=d.day, m=d.month, Y=d.year, dd=f"{d.day:02d}", mm=f"{d.month:02d}")


def day_cell_values(cfg: dict, day: Day) -> dict:
    """Cell -> value for one day sheet (None = leave blank)."""
    c, cr = cfg["cells"], cfg["credit_rows"]
    values = {ref: None for ref in day_input_cells(cfg)}
    values[c["rate_petrol"]] = day.rates["petrol"]
    values[c["rate_diesel"]] = day.rates["diesel"]
    values[c["rate_power"]] = day.rates["power"]
    for key in ("paytm", "hp_card", "icici", "phonepe", "cash", "credit"):
        values[c[key]] = day.payments[key]
    values[c["icici_copy"]] = day.payments["icici"]

    rows = c["nozzle_rows"]
    for fuel in ("petrol", "diesel", "power"):
        col = c[f"{fuel}_col"]
        amounts = [a for _, a in day.nozzles[fuel]]
        if c["fuel_amounts"] == "sum" or len(amounts) > len(rows):
            # keep the first rows one per nozzle, put the rest together in the last row
            head = [] if c["fuel_amounts"] == "sum" else amounts[: len(rows) - 1]
            amounts = head + [round(sum(amounts[len(head):]), 2)]
        for r, amount in zip(rows, amounts):
            values[f"{col}{r}"] = amount

    type_names = cr["type_names"]
    for i, row in enumerate(day.credit_rows):
        r = cr["first_row"] + i
        values[f"{cr['party_col']}{r}"] = row["party"]
        values[f"{cr['type_col']}{r}"] = type_names[row["fuel"]]
        values[f"{cr['amount_col']}{r}"] = row["amount"]
        values[f"{cr['invoice_col']}{r}"] = row["invoice"]
    return values


def _write_register(cfg: dict, ws, days: list[Day]) -> dict:
    """Month credit register (Sheet1): date, party, type, qty, rate, amount, inv, Total row."""
    type_names = cfg["credit_rows"]["type_names"]
    written = {}
    r = cfg["workbook"]["register_first_row"]
    for day in days:
        if not day.credit_rows:
            continue
        start = r
        for row in day.credit_rows:
            vals = {
                "A": datetime(day.date.year, day.date.month, day.date.day),
                "B": row["party"],
                "C": type_names[row["fuel"]],
                "D": f"=+F{r}/E{r}",
                "E": day.rates[row["fuel"]],
                "F": row["amount"],
                "G": row["invoice"],
            }
            for col, v in vals.items():
                ws[f"{col}{r}"] = v
                written[f"{col}{r}"] = v
            ws[f"A{r}"].number_format = "dd-mm-yyyy"
            ws[f"F{r}"].number_format = "0.00"
            r += 1
        ws[f"B{r}"] = "Total"
        ws[f"F{r}"] = f"=SUM(F{start}:F{r - 1})"
        ws[f"F{r}"].number_format = "0.00"
        written[f"F{r}"] = ws[f"F{r}"].value
        r += 1
    return written


def write_month(cfg: dict, days: list[Day], year: int, month: int) -> Path:
    template = cfg["paths"]["template"]
    if not template.exists():
        raise WorkbookError(f"{template} not found - run: python -m skfs_ocr make-template <your sample.xlsx>")
    wb = openpyxl.load_workbook(template)
    tpl_name = cfg["workbook"]["template_day_sheet"]
    if tpl_name not in wb.sheetnames:
        raise WorkbookError(f"sheet {tpl_name!r} not in {template}")
    tpl = wb[tpl_name]

    expected = {}
    days = sorted(days, key=lambda d: d.date)
    for day in days:
        ws = wb.copy_worksheet(tpl)
        ws.title = sheet_name(cfg, day.date)
        values = day_cell_values(cfg, day)
        for ref, v in values.items():
            ws[ref].value = v
        expected[ws.title] = values

    reg_name = cfg["workbook"]["register_sheet"]
    wb.remove(tpl)
    if reg_name and reg_name in wb.sheetnames:
        reg = wb[reg_name]
        expected[reg_name] = _write_register(cfg, reg, days)
        wb.move_sheet(reg, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(reg_name))
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True   # Excel recalculates every formula on open

    out = output_path(cfg, year, month)
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=out.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        _verify(tmp, expected)
        try:
            os.replace(tmp, out)
        except PermissionError:
            raise WorkbookError(f"{out.name} is open in Excel - close it and run again")
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return out


def _verify(path: str, expected: dict) -> None:
    """Re-open the saved file and compare every written cell."""
    wb = openpyxl.load_workbook(path)
    bad = []
    for sheet, cells in expected.items():
        ws = wb[sheet]
        for ref, v in cells.items():
            got = ws[ref].value
            if isinstance(v, float) and isinstance(got, (int, float)):
                ok = abs(got - v) < 1e-9
            else:
                ok = got == v
            if not ok:
                bad.append(f"{sheet}!{ref}: wrote {v!r}, file has {got!r}")
    if bad:
        raise WorkbookError("saved workbook does not match:\n  " + "\n  ".join(bad[:20]))
